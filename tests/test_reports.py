"""Toplantı sonu raporu: DOCX + PDF üretimi ve servis akışı (Claude sahte)."""

from __future__ import annotations

from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from katip.api.main import app
from katip.config import get_settings
from katip.events import Segment
from katip.reports.render import ReportContext, build_html, render_docx, render_pdf
from katip.reports.service import generate_report
from katip.reports.summarize import ActionItem, Decision, MeetingReport, ParticipantNote
from katip.store import repo
from katip.store.org import OrgChart, apply_org_chart

REPORT = MeetingReport(
    title="Q4 Pazarlama Planı",
    summary="Ekip Q4 kampanya bütçesini ve lansman tarihini konuştu.",
    key_points=["Kampanya bütçesi", "Lansman tarihi"],
    decisions=[Decision(text="Lansman 15 Kasım'da yapılacak", owner="Deniz Kaya")],
    action_items=[
        ActionItem(
            owner="Ayşe Yılmaz", task="Bütçe tablosunu hazırla", due="Cuma", priority="yüksek"
        )
    ],
    open_questions=["Ajans seçimi kimde?"],
    participants=[ParticipantNote(name="Ayşe Yılmaz", highlights=["Bütçe önerisi sundu"])],
)

SEGMENTS = [
    {"start_s": 0.0, "speaker_name": "Deniz Kaya", "text": "Q4 planına bakalım."},
    {"start_s": 4.5, "speaker_name": "Ayşe Yılmaz", "text": "Bütçe tablosunu Cuma'ya hazırlarım."},
    {"start_s": 9.0, "speaker_name": "Deniz Kaya", "text": "Lansman 15 Kasım olsun."},
]


async def _real_jpeg() -> bytes:
    """Gerçek Chromium ekran görüntüsü (bot ile aynı yol)."""
    from playwright.async_api import async_playwright

    try:
        async with async_playwright() as pw:
            browser = await pw.chromium.launch(headless=True)
            page = await browser.new_page(viewport={"width": 640, "height": 360})
            await page.set_content("<h1 style='font-family:sans-serif'>Q4 Bütçe</h1>")
            data = await page.screenshot(type="jpeg", quality=80)
            await browser.close()
            return data
    except Exception as exc:
        pytest.skip(f"Chromium yok: {exc}")


def _ctx() -> ReportContext:
    start = datetime(2026, 9, 29, 10, 0, tzinfo=UTC)
    return ReportContext(
        report=REPORT,
        meeting_id="m-test",
        started_at=start,
        ended_at=start + timedelta(minutes=42),
        platform="teams",
        company="Example",
        participants=[("Deniz Kaya", "Kurucu"), ("Ayşe Yılmaz", "Pazarlama Müdürü")],
        segments=SEGMENTS,
    )


def test_docx_contains_all_sections(tmp_path: Path) -> None:
    path = render_docx(_ctx(), tmp_path / "r.docx")
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    for expected in (
        "Q4 Pazarlama Planı",
        "Yönetici özeti",
        "Kararlar",
        "Lansman 15 Kasım'da yapılacak",
        "Açık sorular",
        "Ek: Transkript",
        "Bütçe tablosunu Cuma'ya hazırlarım.",
    ):
        assert expected in text
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert "Ayşe Yılmaz" in cells and "Bütçe tablosunu hazırla" in cells and "Cuma" in cells
    assert "42 dk" in cells and "Deniz Kaya (Kurucu), Ayşe Yılmaz (Pazarlama Müdürü)" in cells


def test_html_escapes_transcript_content() -> None:
    ctx = _ctx()
    ctx.segments = [{"start_s": 0, "speaker_name": "X", "text": "<script>alert(1)</script>"}]
    out = build_html(ctx)
    assert "<script>alert" not in out and "&lt;script&gt;" in out


async def test_pdf_is_rendered(tmp_path: Path) -> None:
    try:
        path = await render_pdf(_ctx(), tmp_path / "r.pdf")
    except Exception as exc:  # Chromium kurulu değilse (CI) atla
        pytest.skip(f"Chromium yok: {exc}")
    data = path.read_bytes()
    assert data.startswith(b"%PDF") and len(data) > 5000


async def test_generate_report_end_to_end(database) -> None:
    await apply_org_chart(
        OrgChart.from_yaml(
            "people:\n  - name: Ayşe Yılmaz\n    title: Pazarlama Müdürü\n    aliases: [Ayşe]\n"
            "  - name: Toplantıya Katılmayan\n    email: secret@example.com\n"
        )
    )
    await repo.upsert_meeting("m1", platform="meet")
    for i, s in enumerate(SEGMENTS):
        await repo.save_segment(
            "m1",
            Segment(
                id=f"s{i}",
                session_id="m1",
                speaker=s["speaker_name"],
                start_s=s["start_s"],
                end_s=s["start_s"] + 3,
                text=s["text"],
                lang="tr",
                is_final=True,
            ),
        )

    seen: dict = {}

    await repo.add_note("m1", "decision", "Lansman 15 Kasım", "Deniz Kaya")
    shot = await _real_jpeg()
    shot_path = Path(get_settings().data_dir) / "shot.jpg"
    shot_path.parent.mkdir(parents=True, exist_ok=True)
    shot_path.write_bytes(shot)
    asset = await repo.add_asset("m1", "screenshot", 1, str(shot_path), 65.0, "Ayşe")
    await repo.set_asset_caption(asset, "Q4 bütçe tablosu: toplam 1,2 milyon TL.")

    async def fake_summarizer(segments, *, org_context, notes, lang):
        seen.update(segments=segments, org=org_context, notes=notes, lang=lang)
        return REPORT

    try:
        docs = await generate_report("m1", summarizer=fake_summarizer)
    except Exception as exc:
        if "Executable doesn't exist" in str(exc):
            pytest.skip("Chromium yok")
        raise

    assert "Ayşe Yılmaz" in seen["org"]
    assert "Toplantıya Katılmayan" not in seen["org"] and "secret@example.com" not in seen["org"]
    assert [s["text"] for s in seen["segments"]] == [s["text"] for s in SEGMENTS]
    assert Path(docs["docx"]).exists() and Path(docs["pdf"]).exists()
    stored = await repo.get_meeting("m1")
    assert stored.report["title"] == "Q4 Pazarlama Planı"
    assert stored.documents == docs

    # Katip'e aldırılan karar ve ekran görüntüsü açıklaması özetleyiciye gitti.
    assert "- Karar: Lansman 15 Kasım (Deniz Kaya)" in seen["notes"]
    assert "- Ekran görüntüsü 1: Q4 bütçe tablosu" in seen["notes"]
    # Word raporunda ekran görüntüsü bölümü ve görselin kendisi var.
    doc = Document(docs["docx"])
    assert any(p.text == "8. Ekran görüntüleri" for p in doc.paragraphs)
    assert any("Görüntü 1 · 01:05" in p.text for p in doc.paragraphs)
    assert len(doc.inline_shapes) == 1


async def test_generate_report_skips_tiny_meetings(database) -> None:
    await repo.upsert_meeting("m2")
    assert await generate_report("m2", summarizer=None) == {}


def test_org_upload_requires_admin_token() -> None:
    client = TestClient(app)
    with client:
        body = "people:\n  - name: Zeynep\n    role: admin\n"
        assert client.put("/api/org", content=body).status_code == 401
        ok = client.put("/api/org", content=body, headers={"Authorization": "Bearer test-admin"})
        assert ok.status_code == 200 and ok.json() == {"people": 1}
        admin = {"Authorization": "Bearer test-admin"}
        assert client.get("/api/people").status_code == 401  # kişi listesi yalnızca yöneticiye
        people = client.get("/api/people", headers=admin).json()
        assert people[0]["name"] == "Zeynep" and people[0]["voice_enrolled"] is False

        assert client.get("/api/meetings/yok/report.pdf").status_code == 401
        missing = client.get("/api/meetings/yok/report.pdf", headers=admin)
        assert missing.status_code == 404


def test_report_date_is_shown_in_local_timezone() -> None:
    """Sunucu UTC; rapor Türkiye saatini göstermeli (canlı testte 06:58 yazıyordu)."""
    ctx = _ctx()
    ctx.started_at = datetime(2026, 9, 30, 6, 58, tzinfo=UTC)
    assert ctx.date_text == "30.09.2026 09:58"
    ctx.started_at = datetime(2026, 9, 30, 6, 58)  # SQLite saat dilimsiz
    assert ctx.date_text == "30.09.2026 09:58"


def test_fallbacks_only_for_models_that_support_them() -> None:
    from katip.reports.summarize import _fallback_args

    # Sonnet 5.5 sunucu tarafı yedeği kabul ediyor (canlı API, 24 Ekim 2026).
    assert _fallback_args("claude-sonnet-5-5")["fallbacks"] == "default"
    assert _fallback_args("claude-haiku-5-5") == {}
    assert _fallback_args("claude-opus-5-5")["fallbacks"] == "default"


def test_thinking_off_uses_the_right_field_per_model() -> None:
    from katip.agent.llm import thinking_off

    # Sonnet 5.5 "disabled"ı reddediyor (canlı API: "send between_tools instead").
    assert thinking_off("claude-sonnet-5-5") == {"thinking": {"type": "between_tools"}}
    assert thinking_off("claude-sonnet-5") == {"thinking": {"type": "disabled"}}
    assert thinking_off("claude-haiku-4-5") == {}
    assert thinking_off("claude-haiku-5-5") == {"thinking": {"type": "disabled"}}


@pytest.mark.parametrize("thinking", [False, True])
async def test_haiku_report_uses_supported_thinking_without_fallbacks(monkeypatch, thinking):
    from types import SimpleNamespace as NS
    from unittest.mock import AsyncMock

    from katip.reports import summarize

    monkeypatch.setenv("KATIP_LLM_MODEL", "claude-haiku-5-5")
    parse = AsyncMock(return_value=NS(stop_reason="end_turn", parsed_output=REPORT, usage=None))
    client = NS(beta=NS(messages=NS(parse=parse)))
    assert await summarize.summarize_meeting(SEGMENTS, client=client, thinking=thinking) == REPORT
    params = parse.call_args.kwargs
    assert params["model"] == "claude-haiku-5-5"
    assert "fallbacks" not in params and "betas" not in params
    assert "output_config" not in params
    if thinking:
        assert "thinking" not in params
    else:
        assert params["thinking"] == {"type": "disabled"}


def test_header_has_topic_times_duration_and_emails() -> None:
    ctx = _ctx()
    ctx.participants = [("Ayşe Yılmaz", "Pazarlama Müdürü", "ayse@example.com"), ("Misafir", None)]
    ctx.ended_at = ctx.started_at + timedelta(minutes=75)
    rows = dict(ctx.meta_rows())
    assert rows["Konu"] == "Q4 Pazarlama Planı"
    assert rows["Tarih"] == "29.09.2026" and rows["Saat"] == "13:00 – 14:15"
    assert rows["Süre"] == "1 sa 15 dk"
    assert rows["Katılımcılar"] == "Ayşe Yılmaz (Pazarlama Müdürü) <ayse@example.com>, Misafir"
    ctx.ended_at = None
    assert dict(ctx.meta_rows())["Saat"] == "13:00 – devam ediyor"


def test_schema_is_fixed_even_when_sections_are_empty() -> None:
    ctx = _ctx()
    ctx.report = MeetingReport(
        title="Kısa toplantı",
        summary="Kısa bir durum toplantısı.",
        key_points=[],
        decisions=[],
        action_items=[],
        open_questions=[],
        participants=[],
    )
    out = build_html(ctx)
    for heading in (
        "1. Yönetici özeti",
        "2. Gündem ve konuşulan başlıklar",
        "3. Kararlar",
        "4. Aksiyonlar",
        "5. Açık sorular",
        "6. Katip&#x27;e aldırılan notlar",
        "7. Katılımcı katkıları",
        "8. Ekran görüntüleri",
    ):
        assert heading in out
    assert out.count("Kayıt yok.") == 7


async def test_final_report_via_batch_api_with_resume() -> None:
    """Batch: yapılandırılmış çıktı şeması isteğe konur, sonuç beklenip ayrıştırılır."""
    from types import SimpleNamespace as NS

    from katip.reports import summarize

    seen: dict = {}
    remembered: list = []

    class FakeBatches:
        polls = 0

        async def create(self, requests):
            seen["params"] = requests[0]["params"]
            return NS(id="msgbatch_1")

        async def retrieve(self, batch_id):
            FakeBatches.polls += 1
            return NS(processing_status="ended" if FakeBatches.polls > 1 else "in_progress")

        async def results(self, batch_id):
            async def gen():
                text = REPORT.model_dump_json()
                usage = NS(
                    input_tokens=1000,
                    output_tokens=100,
                    cache_read_input_tokens=0,
                    cache_creation_input_tokens=0,
                )
                msg = NS(content=[NS(type="text", text=text)], usage=usage)
                yield NS(result=NS(type="succeeded", message=msg))

            return gen()

    client = NS(messages=NS(batches=FakeBatches()))

    async def remember(batch_id):
        remembered.append(batch_id)

    summarize._BATCH_POLL_S = 0
    report = await summarize.summarize_meeting(
        SEGMENTS, client=client, batch=True, on_batch=remember
    )
    assert report.title == "Q4 Pazarlama Planı" and remembered == ["msgbatch_1"]
    fmt = seen["params"]["output_config"]["format"]
    assert fmt["type"] == "json_schema" and "action_items" in str(fmt["schema"])
    # Sunucu yeniden başladıysa var olan batch beklenir, yenisi açılmaz.
    seen.clear()
    await summarize.summarize_meeting(
        SEGMENTS, client=client, batch=True, resume_batch="msgbatch_1"
    )
    assert "params" not in seen


async def test_late_batch_is_cancelled_and_report_made_immediately(monkeypatch) -> None:
    from types import SimpleNamespace as NS

    from katip.reports import summarize

    monkeypatch.setenv("KATIP_REPORT_BATCH_MAX_WAIT_S", "0")
    get_settings.cache_clear()
    cancelled = []

    class SlowBatches:
        async def create(self, requests):
            return NS(id="b1")

        async def retrieve(self, batch_id):
            return NS(processing_status="in_progress")

        async def cancel(self, batch_id):
            cancelled.append(batch_id)

    class Beta:
        class messages:  # noqa: N801
            @staticmethod
            async def parse(**kw):
                usage = NS(
                    input_tokens=1,
                    output_tokens=1,
                    cache_read_input_tokens=0,
                    cache_creation_input_tokens=0,
                )
                return NS(stop_reason="end_turn", parsed_output=REPORT, usage=usage)

    client = NS(messages=NS(batches=SlowBatches()), beta=Beta())
    summarize._BATCH_POLL_S = 0
    report = await summarize.summarize_meeting(SEGMENTS, client=client, batch=True)
    assert report.title == REPORT.title and cancelled == ["b1"]
